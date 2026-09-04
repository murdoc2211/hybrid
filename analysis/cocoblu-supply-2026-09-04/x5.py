import pandas as pd, numpy as np
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter
import openpyxl
df=pd.read_pickle('v7df.pkl'); pe=pd.read_pickle('po_exp.pkl')
pe['Remaining quantity']=pd.to_numeric(pe['Remaining quantity'],errors='coerce').fillna(0)
pe['Cost']=pd.to_numeric(pe['Cost'],errors='coerce').fillna(0)
pe=pe[(pe['Vendor code']=='QZ73J')&(pe['Ship-to location']=='ISK3')]
import datetime
_we=pe.groupby('PO')['Window end'].max()
pe=pe[~pe.PO.isin(_we[_we.dt.date<datetime.date(2026,9,4)].index)]  # drop expired POs
op=pe[pe['Remaining quantity']>0].sort_values(['Order date','PO'])
COVER=35

def _t(r):
    if r.eol: return 'X-EOL (not in Tally FY26-27)'
    if r.soldout: return 'X-No supply'
    if r.arrival=='A-new PO': return 'P0-New arrival (uncapped)'
    if r.arrival=='B-never sold (recent)': return 'P0b-New arrival, never sold'
    if r.franchise!='': return 'P1-Incumbent (running down)'
    if r.drr_sep==0: return 'P5-No demand'
    if r.doh_now<15: return 'P2-Critical'
    if r.ship>0: return 'P3-Top-up to 35d'
    return 'P4-Hold (at/over 35d)'
df['tier']=df.apply(_t,axis=1)

# allocate ship qty back to PO lines, oldest PO first
rows=[]
for asin,g in op.groupby('ASIN',sort=False):
    left=float(df.loc[asin,'ship']) if asin in df.index else 0.0
    for _,r in g.iterrows():
        q=min(left,r['Remaining quantity']); left-=q
        rows.append(dict(PO=r['PO'],Order_date=r['Order date'].date(),ASIN=asin,SKU=r['Model number'],
          Product=str(r['Product name'])[:95],Cost=r['Cost'],Open_qty=int(r['Remaining quantity']),
          SHIP_NOW=int(round(q)),Hold=int(r['Remaining quantity']-round(q)),
          Window_start=r['Window start'].date(),Window_end=r['Window end'].date(),
          Cancel_by=r['Cancellation deadline'].date() if pd.notna(r['Cancellation deadline']) else None))
a=pd.DataFrame(rows); a['Ship_value']=(a.SHIP_NOW*a.Cost).round(0)
a=a.join(df['tier'].rename('Tier'),on='ASIN').sort_values(['Tier','Ship_value'],ascending=[True,False])

sku=df.reset_index().rename(columns={'index':'ASIN'})
sku=sku[['ASIN','sku','name','tier','arrival','franchise','sellable','drr_sep','doh_now','need','open_po','ship','doh_after','hold','gap','cost','ship_val','open_pos','soldout']]
sku.columns=['ASIN','SKU','Product','Tier','New arrival','Franchise','Sellable @3Sep','DRR (1-3 Sep)','DOH now','Need @35d','Open PO','SHIP NOW','DOH after','Hold on PO','Gap - need PO','Vendor cost','Ship value INR','Open POs','No supply']
sku=sku[(sku['Sellable @3Sep']>0)|(sku['Open PO']>0)|(sku['DRR (1-3 Sep)']>0)].sort_values(['Tier','Ship value INR'],ascending=[True,False])

FRG={'GIGA 20000':['B0DMDZF5SV','B0HFJGKC8H'],'Click 10000 magnetic':['B0DFZ3FK9F','B0HGB1DJKY','B0HGB1T2S7','B0CG668622'],'Quad Pro cables (coexist)':['B0D8443PTW','B0H71NCHP7']}
fr=pd.DataFrame([[f,round(df.loc[m,'drr_sep'].sum(),1),int(df.loc[m,'sellable'].sum()),int(df.loc[m,'ship'].sum()),
  int(df.loc[m,'sellable'].sum()+df.loc[m,'ship'].sum()),
  round((df.loc[m,'sellable'].sum()+df.loc[m,'ship'].sum())/df.loc[m,'drr_sep'].sum(),0)]
  for f,mm in FRG.items() for m in [[x for x in mm if x in df.index]]],
  columns=['Franchise','Combined DRR','Stock now','Ship','Total after','Franchise DOH after'])

hold=a[a.Hold>0].copy(); hold['Hold_value']=(hold.Hold*hold.Cost).round(0)
hold=hold[['Tier','PO','Order_date','ASIN','SKU','Product','Open_qty','SHIP_NOW','Hold','Cost','Hold_value','Cancel_by']].sort_values('Hold_value',ascending=False)
gap=sku[sku['Gap - need PO']>0][['ASIN','SKU','Product','Sellable @3Sep','DRR (1-3 Sep)','DOH now','Need @35d','Open PO','Gap - need PO','Vendor cost']].copy()
gap['Gap value INR']=(gap['Gap - need PO']*gap['Vendor cost']).round(0)
posum=a[a.SHIP_NOW>0].groupby(['PO','Order_date','Window_start','Window_end'],as_index=False)[['SHIP_NOW','Ship_value']].sum().sort_values('Ship_value',ascending=False)

sh=df[df.ship>0]
summ=pd.DataFrame([
 ['Scope','Vendor code QZ73J / FC ISK3 only'],
 ['Cover policy','35d cap on ESTABLISHED SKUs only'],
 ['New arrivals','No cap - ship Amazon full ask (no DRR to cap on)'],
 ['Sellable inventory at Cocoblu (3 Sep)',int(df.sellable.sum())],
 ['Open PO in scope (7 POs, ASN = 0)',int(df.open_po.sum())],
 ['RECOMMENDED SHIP NOW (units)',int(df.ship.sum())],
 ['Recommended ship value (INR, vendor cost)',int(df.ship_val.sum())],
 ['Hold / cancel (units)',int(df.hold.sum())],
 ['Units needed with no PO cover (ask Amazon)',int(df.gap.sum())],
 ['SKUs shipped',int(len(sh))],
 ['Max DOH - established SKUs after ship','%.0f'%df[(df.arrival=='')&(df.drr_sep>0)&(df.ship>0)].doh_after.max()],
 ['New arrival units (group A - new PO)',int(df[df.arrival=='A-new PO'].ship.sum())],
 ['New arrival units (group B - never sold)',int(df[df.arrival=='B-never sold (recent)'].ship.sum())],
 ['Excluded - no supply','Click 20000, Quad Pro 1.5m 60W'],
 ['Excluded - EOL (no Tally sale FY26-27)','%d SKUs / %du / Rs %s'%(int((df.eol&(df.open_po>0)).sum()),int(df[df.eol].open_po.sum()),'{:,.0f}'.format((df[df.eol].open_po*df[df.eol].cost).sum()))],
 ['Excluded - other vendor code','PPAFS / 147u at HBA4, HKA2, HNR4, HPN6'],
 ['Excluded - EXPIRED PO','39VRVKCF / 1,105u - window closed 1 Sep, get it cancelled'],
],columns=['Item','Value'])

OUT='/home/user/hybrid/analysis/cocoblu-supply-2026-09-04/Cocoblu_Supply_Plan_v5_ISK3.xlsx'
with pd.ExcelWriter(OUT,engine='openpyxl') as w:
    summ.to_excel(w,sheet_name='0_Summary',index=False)
    posum.to_excel(w,sheet_name='1_Dispatch by PO',index=False)
    a[a.SHIP_NOW>0].to_excel(w,sheet_name='2_Dispatch lines',index=False)
    sku.to_excel(w,sheet_name='3_SKU plan 35d',index=False)
    fr.to_excel(w,sheet_name='4_Franchise cap',index=False)
    gap.to_excel(w,sheet_name='5_Ask Amazon for PO',index=False)
    hold.to_excel(w,sheet_name='6_Hold or cancel',index=False)
    eolsh=df[df.eol&(df.open_po>0)].reset_index().rename(columns={'index':'ASIN'})[['ASIN','sku','name','sellable','open_po','cost','open_pos']]
    eolsh.columns=['ASIN','SKU','Product','Sellable @3Sep','Open PO to cancel','Vendor cost','Open POs']
    eolsh['Cancel value INR']=(eolsh['Open PO to cancel']*eolsh['Vendor cost']).round(0)
    eolsh.to_excel(w,sheet_name='7_EOL - cancel these',index=False)
wb=openpyxl.load_workbook(OUT); fill=PatternFill('solid',fgColor='1F3864')
for ws in wb.worksheets:
    for c in ws[1]: c.font=Font(bold=True,color='FFFFFF'); c.fill=fill; c.alignment=Alignment(wrap_text=True,vertical='center')
    ws.freeze_panes='A2'
    if ws.max_row>1: ws.auto_filter.ref=ws.dimensions
    for i,col in enumerate(ws.iter_cols(min_row=1,max_row=min(ws.max_row,300)),1):
        L=max((len(str(c.value)) for c in col if c.value is not None),default=8)
        ws.column_dimensions[get_column_letter(i)].width=min(max(L+2,9),48)
wb.save(OUT)
print(summ.to_string(index=False)); print(); print(posum.to_string(index=False))
a.to_pickle('a5.pkl')
