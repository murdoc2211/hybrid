import pandas as pd, numpy as np
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter
df=pd.read_pickle('df.pkl'); a=pd.read_pickle('alloc.pkl')
X='/root/.claude/uploads/33e36581-ad8c-550e-954b-50c2cc982611/59cab6dd-Sales_Inv_Open_PO_Stuffcool_04Sep2026.xlsx'
inv=pd.read_excel(X,'inv')
OUT='/home/user/hybrid/Cocoblu_Supply_Plan_04Sep2026.xlsx'

# Sheet: SKU plan
sku=df.reset_index().rename(columns={'index':'ASIN'})
sku=sku[['ASIN','sku','name','tier','sellable','drr_sep','doh_now','gvs','cover_tgt','need_units','open_po','ship_now','hold_po','gap_no_po','doh_after','cost','ship_val','open_pos','oldest_open']]
sku.columns=['ASIN','SKU','Product','Tier','Sellable @3Sep','DRR (1-3 Sep)','DOH now','GV (3d)','Cover target (d)','Need units','Open PO','SHIP NOW','Hold on PO','Gap - need new PO','DOH after ship','Vendor cost','Ship value INR','Open POs','Oldest open PO']
sku=sku[(sku['Sellable @3Sep']>0)|(sku['Open PO']>0)|(sku['DRR (1-3 Sep)']>0)].sort_values(['Tier','Ship value INR'],ascending=[True,False])

disp=a[a.SHIP_NOW>0][['Tier','FC','PO','Order_date','Window_start','Window_end','Cancel_by','ASIN','SKU','Product','Open_qty','SHIP_NOW','Hold','Cost','Ship_value']]
hold=a[a.Hold>0][['Tier','FC','PO','Order_date','ASIN','SKU','Product','Open_qty','SHIP_NOW','Hold','Cost']].assign(Hold_value=lambda d:(d.Hold*d.Cost).round(0)).sort_values('Hold_value',ascending=False)
gap=sku[sku['Gap - need new PO']>0][['ASIN','SKU','Product','Tier','Sellable @3Sep','DRR (1-3 Sep)','DOH now','Need units','Open PO','Gap - need new PO','Vendor cost']].copy()
gap['Gap value INR']=(gap['Gap - need new PO']*gap['Vendor cost']).round(0)

fcsum=a[a.SHIP_NOW>0].groupby(['FC','PO','Order_date','Window_start','Window_end'],as_index=False)[['SHIP_NOW','Ship_value']].sum().sort_values('Ship_value',ascending=False)
invp=inv[inv.disposition.str.upper()=='SELLABLE'].pivot_table(index='asin',columns='age_group',values='qty',aggfunc='sum',fill_value=0)
invp['TOTAL SELLABLE']=invp.sum(axis=1)
invp=invp.join(df['name'].rename('Product')).sort_values('TOTAL SELLABLE',ascending=False).reset_index()

summ=pd.DataFrame([
 ['Sellable inventory at Cocoblu (3 Sep)',int(df.sellable.sum())],
 ['Unsellable inventory',int(inv[inv.disposition.str.upper()=='UNSELLABLE'].qty.sum())],
 ['Total open PO units (11 POs, nothing shipped - ASN = 0)',int(a.Open_qty.sum())],
 ['Recommended SHIP NOW units',int(a.SHIP_NOW.sum())],
 ['Recommended ship value (INR, vendor cost)',int(a.Ship_value.sum())],
 ['Hold / renegotiate units',int(a.Hold.sum())],
 ['Units needed but NOT on any open PO (ask for fresh PO)',int(sku['Gap - need new PO'].sum())],
 ['Account DRR (1-3 Sep, units/day)',round(df.drr_sep.sum(),1)],
 ['Account DOH now',round(df.sellable.sum()/df.drr_sep.sum(),1)],
 ['Account DOH after recommended ship',round((df.sellable.sum()+a.SHIP_NOW.sum())/df.drr_sep.sum(),1)],
],columns=['Metric','Value'])

with pd.ExcelWriter(OUT,engine='openpyxl') as w:
    summ.to_excel(w,sheet_name='0_Summary',index=False)
    fcsum.to_excel(w,sheet_name='1_Dispatch by PO',index=False)
    disp.to_excel(w,sheet_name='2_Dispatch lines',index=False)
    sku.to_excel(w,sheet_name='3_SKU plan',index=False)
    gap.to_excel(w,sheet_name='4_PO gap - ask Amazon',index=False)
    hold.to_excel(w,sheet_name='5_Hold or cancel',index=False)
    invp.to_excel(w,sheet_name='6_Sellable inv pivot',index=False)

import openpyxl
wb=openpyxl.load_workbook(OUT)
fill=PatternFill('solid',fgColor='1F3864')
for ws in wb.worksheets:
    for c in ws[1]: c.font=Font(bold=True,color='FFFFFF'); c.fill=fill; c.alignment=Alignment(wrap_text=True,vertical='center')
    ws.freeze_panes='A2'
    if ws.max_row>1: ws.auto_filter.ref=ws.dimensions
    for i,col in enumerate(ws.iter_cols(min_row=1,max_row=min(ws.max_row,300)),1):
        L=max(len(str(c.value)) for c in col if c.value is not None)
        ws.column_dimensions[get_column_letter(i)].width=min(max(L+2,9),48)
wb.save(OUT)
print('written',OUT)
print(summ.to_string(index=False))
